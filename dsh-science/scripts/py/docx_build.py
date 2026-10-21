#!/usr/bin/env python3
"""把 build.js / docx.js 生成的排版规格渲染成 .docx。

为什么要单独一个 Python 进程：python-docx 已经处理好了 OOXML 的命名空间、
样式继承和关系图；在 JS 里手写这些等于重造一遍而且还容易在 Word/WPS 之间出现
细微差异。这里只吃一个 JSON 文件，不改任何业务逻辑。

输入 JSON（由 lib/science/format/docx.js 生成）：
{
  "out": "绝对路径.docx",
  "spec": {                          # 版式，来自 data/journal-formats.json 的 docx 段
    "pageSize": "A4" | "Letter",
    "marginsMm": {"top":25,"right":25,"bottom":25,"left":25},
    "bodyFont": "Times New Roman", "cjkFont": "宋体",
    "bodyPt": 12, "lineSpacing": 2, "lineNumbers": true,
    "titlePt": 16, "headingPt": 13, "headingBold": true,
    "paragraphSpaceAfterPt": 0
  },
  "blocks": [ ... 见下 ... ]
}

块类型：title / authors / abstract / keywords / heading / para / list /
        figure / table / equation / pagebreak / refs
"""

import json
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PAGE_SIZES = {"A4": (210, 297), "Letter": (215.9, 279.4)}


def set_run_font(run, spec, size=None, bold=None, italic=None):
    latin = spec.get("bodyFont", "Times New Roman")
    east_asia = spec.get("cjkFont", "宋体")
    run.font.name = latin
    run.font.size = Pt(size if size is not None else spec.get("bodyPt", 12))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def add_line_numbers(section, distance_twips=360):
    """连续行号：投稿版几乎都要，Word UI 里藏在「布局 → 行号」。"""
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:lnNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), str(distance_twips))
    sect_pr.append(ln)


def build_footer(section, spec):
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    set_run_font(run, spec, size=spec.get("bodyPt", 12) - 2)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


def add_runs(paragraph, runs, spec, base_bold=None, base_italic=None, size=None):
    for item in runs or []:
        text = item.get("text", "")
        if text == "":
            continue
        run = paragraph.add_run(text)
        set_run_font(
            run,
            spec,
            size=size,
            bold=item.get("bold", base_bold) if item.get("bold") is not None else base_bold,
            italic=item.get("italic", base_italic) if item.get("italic") is not None else base_italic,
        )
        if item.get("code"):
            run.font.name = "Consolas"
            rpr = run._element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                fonts.set(qn("w:ascii"), "Consolas")
                fonts.set(qn("w:hAnsi"), "Consolas")
        if item.get("sup"):
            run.font.superscript = True
        if item.get("url"):
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            run.font.underline = True
    return paragraph


def set_cell_border(cell, edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "bottom", "left", "right"):
        if edge not in edges:
            continue
        spec_value = edges[edge]
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), spec_value.get("val", "single"))
        element.set(qn("w:sz"), str(spec_value.get("sz", 8)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), spec_value.get("color", "000000"))


def build_table(doc, block, spec):
    header = block.get("header") or []
    rows = block.get("rows") or []
    cols = max(len(header), max((len(r) for r in rows), default=0), 1)

    table = doc.add_table(rows=len(rows) + 1, cols=cols)
    table.autofit = True
    table.alignment = 1  # center

    def fill(cell, runs, bold=False):
        cell.paragraphs[0].text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(paragraph, runs, spec, base_bold=bold)
        for run in paragraph.runs:
            set_run_font(run, spec, size=spec.get("bodyPt", 12) - 1, bold=bold)

    for index in range(cols):
        fill(table.rows[0].cells[index], header[index] if index < len(header) else [], bold=True)

    for r, row in enumerate(rows):
        for c in range(cols):
            fill(table.rows[r + 1].cells[c], row[c] if c < len(row) else [], bold=False)

    # 三线表：只在表顶、表头下、表底画线
    last_row = len(rows)
    for c in range(cols):
        set_cell_border(table.rows[0].cells[c], {"top": {"sz": 12}, "bottom": {"sz": 8}})
        set_cell_border(table.rows[last_row].cells[c], {"bottom": {"sz": 12}})

    if block.get("caption"):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(4)
        label = f"Table {block.get('number', '')}. " if block.get("number") else "Table. "
        add_runs(para, [{"text": label, "bold": True}] + block["caption"], spec,
                 size=spec.get("bodyPt", 12) - 1)


def build(doc_spec):
    spec = doc_spec["spec"]
    document = Document()

    # 默认样式
    style = document.styles["Normal"]
    style.font.name = spec.get("bodyFont", "Times New Roman")
    style.font.size = Pt(spec.get("bodyPt", 12))
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), spec.get("bodyFont", "Times New Roman"))
    fonts.set(qn("w:hAnsi"), spec.get("bodyFont", "Times New Roman"))
    fonts.set(qn("w:eastAsia"), spec.get("cjkFont", "宋体"))

    section = document.sections[0]
    width, height = PAGE_SIZES.get(spec.get("pageSize", "A4"), PAGE_SIZES["A4"])
    section.page_width = Mm(width)
    section.page_height = Mm(height)
    margins = spec.get("marginsMm", {})
    section.top_margin = Mm(margins.get("top", 25))
    section.bottom_margin = Mm(margins.get("bottom", 25))
    section.left_margin = Mm(margins.get("left", 25))
    section.right_margin = Mm(margins.get("right", 25))

    if spec.get("lineNumbers"):
        add_line_numbers(section)
    build_footer(section, spec)

    spacing = spec.get("lineSpacing", 1)
    space_after = Pt(spec.get("paragraphSpaceAfterPt", 0))

    def body_para(align=None, size=None, bold=None):
        para = document.add_paragraph()
        fmt = para.paragraph_format
        if spacing >= 2:
            fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        else:
            fmt.line_spacing = spacing
        fmt.space_after = space_after
        fmt.space_before = Pt(0)
        if align is not None:
            para.alignment = align
        return para

    for block in doc_spec["blocks"]:
        kind = block.get("k")

        if kind == "title":
            para = body_para(WD_ALIGN_PARAGRAPH.CENTER)
            add_runs(para, [{"text": block.get("text", ""), "bold": True}], spec,
                     size=spec.get("titlePt", 16))

        elif kind == "authors":
            for entry in block.get("entries", []):
                para = body_para(WD_ALIGN_PARAGRAPH.CENTER)
                add_runs(para, [{"text": entry.get("name", "")}], spec, size=spec.get("bodyPt", 12))
                if entry.get("affiliation"):
                    sub = body_para(WD_ALIGN_PARAGRAPH.CENTER)
                    add_runs(sub, [{"text": entry["affiliation"], "italic": True}], spec,
                             size=spec.get("bodyPt", 12) - 1)
                if entry.get("email"):
                    sub = body_para(WD_ALIGN_PARAGRAPH.CENTER)
                    add_runs(sub, [{"text": entry["email"]}], spec, size=spec.get("bodyPt", 12) - 1)

        elif kind == "abstract":
            para = body_para()
            add_runs(para, [{"text": block.get("label", "Abstract"), "bold": True}], spec)
            para2 = body_para()
            add_runs(para2, block.get("runs", []), spec)

        elif kind == "keywords":
            para = body_para()
            add_runs(para, [{"text": "Keywords: ", "bold": True}] + block.get("runs", []), spec)

        elif kind == "heading":
            level = min(max(block.get("level", 1), 1), 3)
            para = body_para()
            para.paragraph_format.space_before = Pt(10 if level == 1 else 8)
            para.paragraph_format.keep_with_next = True
            size = spec.get("headingPt", 13) if level == 1 else spec.get("bodyPt", 12) + 0.5
            add_runs(para, [{"text": block.get("text", ""), "bold": spec.get("headingBold", True)}],
                     spec, size=size)

        elif kind == "para":
            para = body_para()
            add_runs(para, block.get("runs", []), spec)

        elif kind == "list":
            for item in block.get("items", []):
                para = body_para()
                para.paragraph_format.left_indent = Mm(6 + 6 * item.get("depth", 0))
                para.paragraph_format.first_line_indent = Mm(-6)
                bullet = f"{item.get('index', 1)}. " if block.get("ordered") else "• "
                add_runs(para, [{"text": bullet}] + item.get("runs", []), spec)

        elif kind == "equation":
            para = body_para(WD_ALIGN_PARAGRAPH.CENTER)
            add_runs(para, block.get("runs", []), spec, base_italic=True)

        elif kind == "figure":
            path = block.get("path")
            para = body_para(WD_ALIGN_PARAGRAPH.CENTER)
            if path and os.path.exists(path):
                run = para.add_run()
                width_mm = block.get("widthMm")
                if width_mm:
                    run.add_picture(path, width=Mm(width_mm))
                else:
                    run.add_picture(path)
            else:
                add_runs(para, [{"text": f"[图片缺失：{block.get('src', '')}]", "italic": True}], spec)
            if block.get("caption"):
                cap = body_para(WD_ALIGN_PARAGRAPH.LEFT)
                cap.paragraph_format.space_after = Pt(10)
                label = f"Figure {block.get('number', '')}. " if block.get("number") else "Figure. "
                add_runs(cap, [{"text": label, "bold": True}] + block["caption"], spec,
                         size=spec.get("bodyPt", 12) - 1)

        elif kind == "table":
            build_table(document, block, spec)

        elif kind == "pagebreak":
            document.add_page_break()

        elif kind == "refs":
            refs = block.get("entries", [])
            for entry in refs:
                para = body_para()
                fmt = para.paragraph_format
                fmt.left_indent = Mm(8)
                fmt.first_line_indent = Mm(-8)
                prefix = f"{entry['n']}. " if entry.get("n") else ""
                add_runs(para, [{"text": prefix + entry.get("text", "")}], spec)

        elif kind == "spacer":
            body_para()

    out = doc_spec["out"]
    os.makedirs(os.path.dirname(out), exist_ok=True)
    document.save(out)
    return out


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"ok": False, "error": "usage: docx_build.py <spec.json>"}))
        return 2
    with open(sys.argv[1], "r", encoding="utf-8") as handle:
        doc_spec = json.load(handle)
    try:
        out = build(doc_spec)
        print(json.dumps({"ok": True, "path": out, "bytes": os.path.getsize(out)}, ensure_ascii=False))
        return 0
    except Exception as exc:  # noqa: BLE001 - 需要把原始错误原样交回 JS 层
        import traceback
        print(json.dumps({
            "ok": False,
            "error": f"{type(exc).__name__}: {exc}",
            "trace": traceback.format_exc()[-1500:],
        }, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())
